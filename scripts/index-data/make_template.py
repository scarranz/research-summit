"""Builds the empty "Summit Index Data" workbook: the layout in docs/DATA_CONTRACT.md.

  python scripts/index-data/make_template.py            # writes templates/Summit Index Data - TEMPLATE.xlsx

Headers only, no data. The analyst pastes the Bloomberg output into it; the (future) extractor
reads every sheet by its header names, never by fixed positions.
"""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(ROOT, "scripts", "spx"))
import gics

OUT = os.path.join(ROOT, "templates", "Summit Index Data - TEMPLATE.xlsx")

# GICS sector -> (S&P 500 sector index, SPDR sector ETF), as in docs/SPX_MARKET_ANALYSIS.md
SECTOR_TICKERS = {
    "Information Technology": ("S5INFT Index", "XLK US Equity"),
    "Communication Services": ("S5TELS Index", "XLC US Equity"),
    "Financials": ("S5FINL Index", "XLF US Equity"),
    "Consumer Discretionary": ("S5COND Index", "XLY US Equity"),
    "Industrials": ("S5INDU Index", "XLI US Equity"),
    "Health Care": ("S5HLTH Index", "XLV US Equity"),
    "Utilities": ("S5UTIL Index", "XLU US Equity"),
    "Consumer Staples": ("S5CONS Index", "XLP US Equity"),
    "Real Estate": ("S5RLST Index", "XLRE US Equity"),
    "Energy": ("S5ENRS Index", "XLE US Equity"),
    "Materials": ("S5MATR Index", "XLB US Equity"),
}
SECTORS = list(gics.SECTORS.values())
IGS = list(gics.INDUSTRY_GROUPS.values())

README = [
    "Summit Index Data — layout contract: docs/DATA_CONTRACT.md",
    "",
    "RAW_SPX / RAW_NDX   Bloomberg's own output. Per month-end: the date in the header cell, constituents below it, their weight (%) in the column to the right. Paste new months to the left; order does not matter.",
    "UNIVERSE            one row per raw ticker ever seen: substitution (ticker), consensus helper (eps_ticker, eps_scale), in_universe flag, labels.",
    "LEVELS              month-end levels. Row 1 = Bloomberg ticker, row 2 = GICS name. SPX, SPY, 11 S5 sector indices, 11 SPDR ETFs, 25 industry-group indices.",
    "W_SECTOR            SPY weights by GICS sector (% , one column per sector + Residual), month-ends.",
    "W_IG                SPY weights by GICS industry group (%), month-ends.",
    "CONSENSUS           long format: date | ticker | metric | period | value.",
    "PRICES              RESERVED for Systems: month-end close, security x date. Not filled by the analyst.",
    "",
    "Dates: month-ends from 2016-12-31. Weights in percent. Blank weight = not in the index that month.",
]


def sheet(wb, name, rows, widths=None):
    ws = wb.create_sheet(name)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF5")
    ws.freeze_panes = "B2" if name not in ("UNIVERSE", "CONSENSUS", "README") else "A2"
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w
    return ws


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet(wb, "README", [[l] for l in README], {"A": 140})
    sheet(wb, "RAW_SPX", [["DATE →  (month-end in the header cell; ticker below it; weight % to its right)"]])
    sheet(wb, "RAW_NDX", [["DATE →  (month-end in the header cell; ticker below it; weight % to its right)"]])
    sheet(wb, "UNIVERSE", [["raw_ticker", "ticker", "composite_id", "eps_ticker", "eps_scale", "in_universe", "name", "sector", "industry_group", "note"]],
          {"A": 22, "B": 22, "C": 18, "D": 22, "E": 10, "F": 12, "G": 36, "H": 26, "I": 40, "J": 40})
    tick = ["DATE", "SPX Index", "SPY US Equity"] + [v[0] for v in SECTOR_TICKERS.values()] + [v[1] for v in SECTOR_TICKERS.values()] + ["" for _ in IGS]
    name = ["", "S&P 500", "SPY"] + list(SECTOR_TICKERS) + list(SECTOR_TICKERS) + IGS
    sheet(wb, "LEVELS", [tick, name])
    sheet(wb, "W_SECTOR", [["DATE"] + SECTORS + ["Residual"]])
    sheet(wb, "W_IG", [["DATE"] + IGS])
    sheet(wb, "CONSENSUS", [["date", "ticker", "metric", "period", "value"]], {"A": 12, "B": 22, "C": 18, "D": 10, "E": 14})
    sheet(wb, "PRICES", [["RESERVED for Systems — month-end close, security rows x month-end date columns. Not filled by the analyst."]])
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
