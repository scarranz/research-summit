import csv,re,json,sys,datetime,warnings,math
warnings.filterwarnings('ignore')
import openpyxl
import os
HERE=os.path.dirname(os.path.abspath(__file__))
S=os.path.join(HERE,"out")+os.sep
os.makedirs(S,exist_ok=True)
rows=list(csv.reader(open(r"G:\My Drive\Summit\Docs\0\BBG_CONSENSUS.txt",encoding='utf-8',errors='replace'),delimiter='\t'))
h=rows[0]
dt=lambda s: datetime.datetime.strptime(s,'%m/%d/%Y')
by={}
for r in rows[1:]:
    if r and r[0]=='MA US EQUITY':
        d=dict(zip(h,r)); k=dt(d['data_as_of'])
        if k not in by or d.get('time','')>by[k].get('time',''): by[k]=d
snaps=[by[k] for k in sorted(by)]
L=snaps[-1]
qn=lambda l:(lambda m:int(m.group(1))*4+int(m.group(2))-1)(re.match(r'(\d{4}) Q(\d)',l))
# ---- integrity: labels + close dates
bad=0
for d in snaps:
    b=qn(d['fq0'])
    for k,o in(('fq-3',-3),('fq+1',1),('fq+2',2),('fq+3',3),('fq+4',4)):
        if qn(d[k])!=b+o: bad+=1
    for o in range(0,5):
        c=d.get('close_fq%s'%('+%d'%o if o else '0'))
        if c:
            m=dt(c); n=b+o
            if (m.year*4+(m.month-1)//3)!=n: bad+=1
print('integrity mismatches',bad)
def val(v,scale):
    if v is None: return None
    v=v.strip().rstrip('.')
    if v=='' or '/' in v or ':' in v or 'Error' in v or '#' in v: return None
    try:x=float(v)
    except:return None
    if scale=='M' and abs(x)<0.001: return None
    return x
# (pref, key, ceName, label, short, group, kind, flip)
C=[
 ('metric1','rev','Net revenue','Net revenue','Net revenue','Totals','head',1),
 ('metric12','opinc','Operating income','Operating income','Op. income','Profit','head',1),
 ('metric13','ebitda','EBITDA','Adjusted EBITDA','EBITDA','Profit','head',1),
 ('metric25','eps','EPS','EPS (GAAP-comparable, diluted)','EPS','Profit','head',1),
 ('metric23','ni','Net income','Net income (GAAP-comparable)','Net income','Profit','head',1),
 ('metric19','pti','Pre-tax income','Pre-tax income','Pre-tax income','Profit','head',1),
 ('metric4','opex','Total operating expenses','Total operating expenses','Opex','Operating costs','head',1),
 ('metric5','ga','G&A expense','General & administrative expense','G&A','Operating costs','head',1),
 ('metric6','am','Advertising & marketing','Advertising & marketing','Adv. & mktg','Operating costs','head',1),
 ('metric26','dna','D&A','Depreciation & amortization','D&A','Operating costs','head',1),
 ('metric29','sbc','Stock-based comp','Stock-based compensation','SBC','Operating costs','head',1),
 ('kpi12','lit','Provision for litigation','Provision for litigation','Litigation','Operating costs','kpi',1),
 ('metric16','intexp','Interest expense','Interest expense','Interest exp.','Below the line','head',1),
 ('metric15','netint','Net interest expense','Net interest expense','Net interest','Below the line','head',1),
 ('metric14','intnet','Interest expense / (income), net','Interest expense / (income), net','Int. exp./(inc.) net','Below the line','head',1),
 ('metric17','invinc','Investment income','Investment income','Invest. income','Below the line','head',1),
 ('metric18','oie','Other income / (expense), net','Other income / (expense), net','Other inc./(exp.)','Below the line','head',1),
 ('metric20','tax','Income tax expense','Income tax expense','Tax expense','Below the line','head',1),
 ('metric36','cfo','Operating cash flow','Operating cash flow','Op. cash flow','Cash flow & capital','head',1),
 ('metric37','fcf','Free cash flow','Free cash flow','FCF','Cash flow & capital','head',1),
 ('metric35','capex','Capex','Capex (headline capex)','Capex','Cash flow & capital','head',-1),
 ('metric39','buyback','Share repurchases','Share repurchases','Buybacks','Cash flow & capital','head',-1),
 ('metric24','shares','Shares outstanding','Diluted weighted-avg. shares','Diluted shares','Cash flow & capital','head',1),
 ('kpi10','amort','Customer & merchant incentive amortization','Amortization of customer & merchant incentives','Incentive amort.','Cash flow & capital','kpi',1),
 ('kpi4','gross','Gross revenue','Gross revenue','Gross revenue','Revenue build (gross to net)','kpi',1),
 ('kpi5','dom','Domestic assessments','Domestic assessments','Domestic assess.','Revenue build (gross to net)','kpi',1),
 ('kpi6','xb','Cross-border volume fees','Cross-border volume fees','Cross-border fees','Revenue build (gross to net)','kpi',1),
 ('kpi7','tp','Transaction processing revenue','Transaction processing revenue','Txn processing','Revenue build (gross to net)','kpi',1),
 ('kpi8','oth','Other revenue','Other revenue','Other revenue','Revenue build (gross to net)','kpi',1),
 ('kpi9','cinc','Client incentives','Client incentives (contra-revenue)','Client incentives','Revenue build (gross to net)','kpi',1),
 ('kpi3','pnr','Payment network revenue','Payment Network (net revenue)','Payment Network','Net revenue by segment','kpi',1),
 ('kpi11','vas','VAS revenue','Value-Added Services & Solutions (net revenue)','VAS','Net revenue by segment','kpi',1),
 ('kpi1','pvol','Purchase volume','Purchase volume (network spend)','Purchase vol','Volumes & transactions','kpi',1),
 ('kpi2','ptx','Switched transactions','Processed (switched) transactions','Processed txns','Volumes & transactions','kpi',1),
]
def meta(pref):
    if pref.startswith('metric'):
        i=pref[6:]; g=lambda f:(L.get(f+i) or '').strip()
    else:
        i=pref[3:]; g=lambda f:(L.get(f+'_kpi'+i) or '').strip()
    return dict(name=g('metric'),code=g('code'),unit=g('unit'),scale=g('scale'))
Qn=sorted({qn(d[k]) for d in snaps for k in('fq-3','fq0','fq+1','fq+2','fq+3','fq+4')})
Y=list(range(2020,2030))
ql=lambda n:'%dQ%s'%(n%4+1,str(n//4)[2:])
qlong=lambda n:'Q%d %d'%(n%4+1,n//4)
data={}
for pref,key,ce_,label,short,grp,kind,flip in C:
    mt=meta(pref); sc=mt['scale']
    act={n:None for n in Qn}; cons={n:None for n in Qn}; ya={y:None for y in Y}; yc={y:None for y in Y}
    hz={n:[None]*4 for n in Qn}
    for d in snaps:
        b=qn(d['fq0']); fy=int(d['fy0'][:4])
        for k,o in(('fq-3',-3),('fq0',0)):
            v=val(d.get('%s_%s'%(pref,k)),sc)
            if v is not None: act[b+o]=v*flip
        v=val(d.get('%s_fy0'%pref),sc)
        if v is not None: ya[fy]=v*flip
        for o in range(1,5):
            v=val(d.get('%s_fq+%d'%(pref,o)),sc)
            if v is not None and (b+o) in hz: hz[b+o][4-o]=v*flip
        v=val(d.get('%s_fy+1'%pref),sc)
        if v is not None and fy+1 in yc: yc[fy+1]=v*flip
    fy0=int(L['fy0'][:4])
    for o in range(1,6):
        v=val(L.get('%s_fy+%d'%(pref,o)),sc)
        if v is not None and fy0+o in yc: yc[fy0+o]=v*flip
    for n in Qn: cons[n]=hz[n][3]
    b=qn(L['fq0'])
    for o in range(1,5):
        v=val(L.get('%s_fq+%d'%(pref,o)),sc)
        if v is not None: cons[b+o]=v*flip
    data[key]=dict(pref=pref,ce=ce_,label=label,short=short,grp=grp,kind=kind,flip=flip,mt=mt,act=act,cons=cons,hz=hz,ya=ya,yc=yc)
# ---- Summit annual (local xlsm Projection History)
wb=openpyxl.load_workbook(r"G:\My Drive\Summit\Docs\Research\DCF\Payments - All\MA\DCF MA - New.xlsm",data_only=True,read_only=True)
R=list(wb['Projection History'].iter_rows(values_only=True)); hdr=R[0]
ycol={hdr[c]:c for c in range(8,36) if isinstance(hdr[c],int)}
code2row={R[i][1]:i for i in range(1,13)}
SM={'rev':'REV','opinc':'OP_INCOME','ebitda':'EBITDA','ni':'EARNINGS','shares':'SHARES','dna':'DNA','cfo':'CFO','capex':'CAPEX','fcf':'FCF','intexp':'INT_EXP'}
summit={}
for key,code in SM.items():
    row=R[code2row[code]]; arr=[]
    for y in Y:
        v=row[ycol[y]]
        if not isinstance(v,(int,float)) or v==0 or y<2022: arr.append(None)
        else: arr.append(int(math.floor(abs(v)+0.5)) if key!='shares' else int(math.floor(v+0.5)))
    summit[key]=arr
# ---- Summit model operating build (same workbook, tab Segments; annual 2018-2029, model actuals through 2025 tie to the Bloomberg actuals)
SEGR=list(wb['Segments'].iter_rows(max_col=20,values_only=True))
SEGCOL={2029:5,2028:6,2027:7,2026:8,2025:9,2024:10,2023:11,2022:12}
def segrow(code=None,label=None):
    for r in SEGR:
        if code is not None and r[3]==code: return r
        if label is not None and r[2]==label: return r
    raise SystemExit('Segments row not found: %s%s'%(code,label))
SEG={'pnr':('SEG1_REV',None,1,False),'vas':('SEG2_REV',None,1,False),'gross':('SEG3_REV',None,1,True),'tp':('SEG6_REV',None,1,False),
     'oth':('SEG7_REV',None,1,True),'cinc':('SEG8_REV',None,1,False),'opex':('OPEX',None,1,False),'ga':('G&A',None,1,False),'am':('MARKETING',None,1,False),
     'sbc':('SEG_SBC',None,1,False),'lit':('PROV_LIT',None,1,False),'tax':('SEG_INCTAX',None,1,False),'buyback':('STOCK_BUYBACK',None,-1,False),
     'ptx':('SWITCH_TRANS',None,1,False),'pvol':(None,'Purchase Volume',1000,False),'eps':('ADJ_EPS',None,1,False)}
for key,(code,label,mult,skip22) in SEG.items():
    r=segrow(code,label); arr=[]
    for y in Y:
        v=r[SEGCOL[y]] if y in SEGCOL else None
        if not isinstance(v,(int,float)) or y<2022 or (skip22 and y==2022): arr.append(None)
        elif key=='eps': arr.append(round(v,2))
        else: arr.append(int(math.floor(abs(v*mult) if key=='buyback' else v*mult+0.5)))
    summit[key]=arr
# ---- Summit QUARTERLY (tab Projection History, quarterly columns, snapshot ID 3 = rows 2-55: the newest save, after the 2Q26 print).
# Columns 1Q26 and 2Q26 hold the model's own pre-print estimates for those quarters (the model was built in 2026, so earlier quarters carry
# only the actuals tab); 3Q26 and 4Q26 are forecasts. Read through 4Q26 ONLY: from 1Q27 the model's quarterly chain breaks on the lines that sum
# other lines (net revenue, payment network, gross, transaction processing, operating income, EBITDA, EBT, tax, net income and SBC go negative or
# explode) and the stand-alone lines do not tie to the model's own 2027 year either (Card related activity: four 2027 quarters sum to 24.5B against 29.0B
# in the year), so no 2027 quarter is loaded.
SQ={'rev':('SEG_REV',1,True),'pnr':('SEG1_REV',1,True),'vas':('SEG2_REV',1,False),'gross':('SEG3_REV',1,True),'card':('SEG4_REV',1,False),
    'tp':('SEG6_REV',1,True),'oth':('SEG7_REV',1,False),'cinc':('SEG8_REV',1,False),'opex':('OPEX',1,False),'ga':('G&A',1,False),
    'am':('MARKETING',1,False),'lit':('PROV_LIT',1,False),'opinc':('SEG_OPINC',1,True),'ebitda':('SEG_EBITDA',1,True),'pti':('SEG_EBT',1,True),
    'tax':('SEG_INCTAX',1,True),'ni':('SEG_NETINC',1,True),'sbc':('SEG_SBC',1,True),'intexp':('SEG_INT_EXP',1,False),
    'invinc':('SEG_INT_INC',-1,False),'intnet':('SEG_NON_OPINC',-1,False),'buyback':('STOCK_BUYBACK',-1,False)}
_hq={hdr[c]:c for c in range(len(hdr)) if isinstance(hdr[c],str) and hdr[c].startswith('Q') and ' - ' in hdr[c]}
_rowq={R[i][1]:i for i in range(1,56) if R[i][1]}
def _qlab(h): a,b=h.replace('Q','').split(' - '); return '%sQ%s'%(a,b[2:])
SQV={}   # key -> {period label: value}, summit's own sign convention already flipped to the dataset's
SQ_PERIODS=['1Q26','2Q26','3Q26','4Q26']
for key,(code,sg,cut) in SQ.items():
    row=R[_rowq[code]]; d={}
    for per in SQ_PERIODS:
        h='Q%s - 20%s'%(per[0],per[2:]); v=row[_hq[h]]
        if isinstance(v,(int,float)) and v!=0: d[per]=v*sg
    SQV[key]=d
_ii=SQV.get('invinc',{}); _ie={}
_ri=R[_rowq['SEG_INT_EXP']]
for per in SQ_PERIODS:
    v=_ri[_hq['Q%s - 20%s'%(per[0],per[2:])]]
    if isinstance(v,(int,float)) and v!=0 and per in _ii: _ie[per]=v+_ii[per]   # interest expense less interest income = net interest expense
SQV['netint']=_ie
def rnd(x,dp):
    if x is None: return None
    f=10**dp; r=math.floor(abs(x)*f+0.5)/f
    r=r if x>=0 else -r
    return int(r) if dp==0 else r
def fnum(v):
    if v is None: return 'null'
    if isinstance(v,int): return str(v)
    s=('%.4f'%v).rstrip('0').rstrip('.')
    return s
js=lambda a:'['+','.join(fnum(v) for v in a)+']'
def unitOf(key):
    if key=='eps': return 'eps'
    if key in('shares','ptx'): return 'cnt'
    return 'usdM'
def ceunit(key):
    if key=='eps': return '$'
    if key in('shares','ptx'): return 'B'
    ar=[abs(v) for v in data[key]['act'].values() if v is not None][-8:]
    med=sorted(ar)[len(ar)//2] if ar else 0
    return '$B' if med>=1000 else '$M'
dp_for=lambda key:2 if key=='eps' else 0
out=[]
out.append("var Q=["+",".join("'%s'"%ql(n) for n in Qn)+"];")
out.append("var Y=["+",".join("'%d'"%y for y in Y)+"];\n")
def block(name,getA,getC,withsum):
    o=["var %s = {"%name]
    for idx,(pref,key,ce_,label,short,grp,kind,flip) in enumerate(C):
        D=data[key]; u=unitOf(key); dp=dp_for(key)
        a=[rnd(v,dp) for v in getA(D)]; c=[rnd(v,dp) for v in getC(D)]
        extra=''
        if key in('opinc','ebitda'): extra=", marginOf:'rev', marginLabel:'%s'"%('operating margin' if key=='opinc' else 'EBITDA margin')
        cm='   // file carries it negative; stored as positive magnitude' if flip==-1 else ''
        s=''
        if withsum and key in summit: s=",\n    summit:"+js(summit[key])
        o.append("  %s: { label:'%s', short:'%s', group:'%s', unit:'%s'%s,%s\n    act:  %s,\n    cons: %s%s }%s"%(key,label.replace("'","\\'"),short.replace("'","\\'"),grp,u,extra,cm,js(a),js(c),s,',' if idx<len(C)-1 else ''))
    o.append("};\n")
    return "\n".join(o)
out.append(block('MA_Q',lambda D:[D['act'][n] for n in Qn],lambda D:[D['cons'][n] for n in Qn],False))
out.append(block('MA_Y',lambda D:[D['ya'][y] for y in Y],lambda D:[D['yc'][y] for y in Y],True))
open(S+'ma_js_block.txt','w',encoding='utf-8').write("\n".join(out))
asOf=[dt(d['data_as_of']).strftime('%Y-%m-%d') for d in snaps]
def r3(x,u):
    if x is None: return None
    if u in('$B','B'): return round(x/1000,3)
    if u=='$': return round(x,2)
    return round(x,1)
ce=["var CE_CONS = {",
 "  src:'Bloomberg (BST) \u00b7 BBG_CONSENSUS.txt snapshot archive',",
 "  asOf:["+",".join("'%s'"%a for a in asOf)+"],",
 "  q:["+",".join("'%s'"%qlong(n) for n in Qn)+"],",
 "  hz:['4q out','3q out','2q out','1q out'],",
 "  nHead:%d,"%sum(1 for c in C if c[6]=='head'),
 "  m:["]
order=[c for c in C if c[6]=='head']+[c for c in C if c[6]=='kpi']
fl=lambda a:'['+','.join('null' if v is None else repr(v) for v in a)+']'
for idx,(pref,key,cen,label,short,grp,kind,flip) in enumerate(order):
    D=data[key]; u=ceunit(key); m=D['mt']
    tags='ok'
    hasact=any(v is not None for v in D['act'].values()); hascons=any(v is not None for v in D['cons'].values())
    if hasact and not hascons: tags='nocons'
    if hascons and not hasact: tags='noact'
    qa=[r3(D['act'][n],u) for n in Qn]
    qy=[r3(D['act'].get(n-4),u) for n in Qn]; qq=[r3(D['act'].get(n-1),u) for n in Qn]
    qr=[[r3(v,u) for v in D['hz'][n]] for n in Qn]
    cm='   // file carries it negative; shown as a positive magnitude' if flip==-1 else ''
    ce.append("    { k:'%s', u:'%s', t:'%s', code:'%s',%s"%(cen.replace("'","\\'"),u,tags,m['code'],cm))
    ce.append("      qr:["+",".join(fl(x) for x in qr)+"],")
    ce.append("      qa:"+fl(qa)+",")
    ce.append("      qy:"+fl(qy)+",")
    ce.append("      qq:"+fl(qq)+" }"+(',' if idx<len(order)-1 else ''))
ce+=["  ]","};"]
open(S+'ce_cons_block.txt','w',encoding='utf-8').write("\n".join(ce))
# legacy CE_ANNUAL (3 metrics, $B)
ann=[]
for key,name in(('rev','Net revenue'),('opinc','Operating income'),('ebitda','EBITDA')):
    D=data[key]; ys=[2022,2023,2024,2025,2026,2027]
    act=[round(D['ya'][y]/1000,1) if D['ya'][y] is not None else None for y in ys]
    bb=[round(D['yc'][y]/1000,1) if D['yc'][y] is not None else None for y in ys]
    sm=[round(summit[key][Y.index(y)]/1000,1) if summit[key][Y.index(y)] is not None else None for y in ys]
    ann.append("    { k:'%s', u:'$B', actual:%s, bbg:%s, summit:%s }"%(name,fl(act),fl(bb),fl(sm)))
open(S+'ce_annual_block.txt','w',encoding='utf-8').write(",\n".join(ann))
json.dump(dict(snaps=asOf,Q=[qlong(n) for n in Qn]),open(S+'meta.json','w'))
print('quarters',ql(Qn[0]),ql(Qn[-1]),len(Qn),'| head',sum(1 for c in C if c[6]=='head'),'kpi',sum(1 for c in C if c[6]=='kpi'))
for k,v in data.items():
    la=[x for x in v['act'].values() if x is not None]
    print(k,'|',v['mt']['name'],'|',ceunit(k),'| n_act',len(la),'last',la[-1:] ,'| n_cons',sum(1 for x in v['cons'].values() if x is not None))
